"""The Mousy sample pack is reproducible, up to date, and its answer key matches the transcripts."""

import csv
import json
import re
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tools"))

import make_mousy_pack  # noqa: E402

PACK = make_mousy_pack.DEFAULT_OUT
TURN_LINE = re.compile(r"^\[(\d{2}):(\d{2}):(\d{2})\] (Mousy Assist|Customer): .+$")


def all_files(folder: Path) -> dict[str, bytes]:
    return {p.relative_to(folder).as_posix(): p.read_bytes() for p in sorted(folder.rglob("*")) if p.is_file()}


@pytest.fixture(scope="module")
def fresh_pack(tmp_path_factory: pytest.TempPathFactory) -> Path:
    out = tmp_path_factory.mktemp("pack")
    make_mousy_pack.build_pack(out)
    return out


@pytest.fixture(scope="module")
def rows() -> list[dict]:
    with open(PACK / "mousy_conversations.csv", encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


@pytest.fixture(scope="module")
def answer_key() -> dict:
    return json.loads((PACK / "answer_key.json").read_text(encoding="utf-8"))


def test_generator_is_reproducible(fresh_pack: Path, tmp_path: Path) -> None:
    make_mousy_pack.build_pack(tmp_path / "again")
    assert all_files(fresh_pack) == all_files(tmp_path / "again")


def test_committed_pack_is_up_to_date(fresh_pack: Path) -> None:
    # If this fails: run `python tools/make_mousy_pack.py` and look at what changed.
    assert all_files(PACK) == all_files(fresh_pack)


def test_forty_conversations_with_valid_ids(rows: list[dict]) -> None:
    assert len(rows) == 40
    assert [r["conversation_id"] for r in rows] == [f"CNV-{n:06d}" for n in range(1, 41)]
    assert {r["channel"] for r in rows} == {"voice", "chat"}


def test_transcripts_use_the_open_format(rows: list[dict]) -> None:
    for row in rows:
        lines = row["transcript"].split("\n")
        assert len(lines) == int(row["turns"])
        assert all(TURN_LINE.match(line) for line in lines), row["conversation_id"]
        assert lines[0].split("] ", 1)[1].startswith("Mousy Assist:"), "the bot always opens"


def test_only_invented_identifiers(rows: list[dict]) -> None:
    text = "\n".join(r["transcript"] for r in rows)
    assert all(1000 <= int(n) <= 9999 for n in re.findall(r"MO-(\d+)", text))
    assert all(p.startswith("555-01") for p in re.findall(r"\b\d{3}-\d{4}\b", text))
    assert all(r["started_at"].startswith("2025-03-") and r["started_at"].endswith("Z") for r in rows)


def test_answer_key_turns_point_at_bot_turns(rows: list[dict], answer_key: dict) -> None:
    by_id = {r["conversation_id"]: r["transcript"].split("\n") for r in rows}
    for entry in answer_key["conversations"]:
        lines = by_id[entry["conversation_id"]]
        for bug in entry["bugs"]:
            assert bug["type"] in answer_key["bug_types"]
            assert "] Mousy Assist: " in lines[bug["turn"] - 1], (entry["conversation_id"], bug)
        first = entry["first_customer_utterance"]
        assert lines[first["turn"] - 1].endswith(f"Customer: {first['text']}")


def test_every_bug_type_is_planted(answer_key: dict) -> None:
    planted = {bug["type"] for entry in answer_key["conversations"] for bug in entry["bugs"]}
    assert planted == set(answer_key["bug_types"])


def test_core_transcript_files_match_their_rows(rows: list[dict], answer_key: dict) -> None:
    by_id = {r["conversation_id"]: r for r in rows}
    core = [e for e in answer_key["conversations"] if e["transcript_file"]]
    assert len(core) == 8
    for entry in core:
        text = (PACK / entry["transcript_file"]).read_text(encoding="utf-8")
        assert text == by_id[entry["conversation_id"]]["transcript"] + "\n"
        golden = PACK / entry["transcript_file"].replace("transcripts/", "golden/").replace(".txt", ".json")
        assert json.loads(golden.read_text(encoding="utf-8"))["conversation_id"] == entry["conversation_id"]


def test_excel_files_match_the_csv(rows: list[dict]) -> None:
    sheet = load_workbook(PACK / "mousy_conversations.xlsx").active
    values = list(sheet.iter_rows(values_only=True))
    assert list(values[0]) == make_mousy_pack.COLUMNS
    assert len(values) == 41
    assert values[1][-1] == rows[0]["transcript"]
    assert load_workbook(PACK / "mousy_batch_20.xlsx").active.max_row == 21
    intents = load_workbook(PACK / "intents_v3.xlsx").active
    assert intents.max_row == 11 and intents.max_column == 7


def test_manifest_lists_real_files_with_true_sizes() -> None:
    manifest = json.loads((PACK / "manifest.json").read_text(encoding="utf-8"))
    for entry in manifest["files"]:
        path = PACK / entry["path"]
        assert path.is_file(), entry["path"]
        assert path.stat().st_size == entry["bytes"]
