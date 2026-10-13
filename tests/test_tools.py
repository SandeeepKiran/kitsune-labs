"""Behaviour of the nine tools (M3 no-AI tools, M4 AI tools).

Every tool must work with "Load Mousy sample" and no key. The AI path is tested against a
mocked provider (Playwright routing), so no real key and no network call is ever used.
"""

import io
import json
import re
import zipfile
from pathlib import Path
from urllib.parse import urlparse

import openpyxl
import pytest
from playwright.sync_api import Browser, Page, Route, expect

ROOT = Path(__file__).resolve().parent.parent
ANSWER_KEY = json.loads((ROOT / "app" / "public" / "sample-data" / "answer_key.json").read_text(encoding="utf-8"))
FAKE_KEY = "test-key-" + "y" * 32   # shape-valid, never a real key
GEMINI_HOST = "generativelanguage.googleapis.com"

BENTO_TOOLS = [
    "bug-finder", "test-case-generator", "utterance-generator", "first-utterance-finder",
    "transcript-converter", "metrics-extractor", "report-builder", "email-templates",
]


def load_sample(page: Page, site_url: str, slug: str, sample: str | None = None) -> None:
    page.goto(f"{site_url}tools/{slug}/" + (f"?sample={sample}" if sample else ""))
    if not sample:
        page.get_by_test_id("load-sample").click()
    expect(page.get_by_test_id("run-status")).to_have_text("Done")


def download(page: Page, button: str | re.Pattern) -> tuple[str, bytes]:
    with page.expect_download() as info:
        page.get_by_role("button", name=button, exact=isinstance(button, str)).click()
    return info.value.suggested_filename, Path(info.value.path()).read_bytes()


def add_key(page: Page) -> None:
    page.get_by_test_id("key-chip").click()
    drawer = page.get_by_test_id("key-drawer")
    drawer.get_by_label("API key").fill(FAKE_KEY)
    drawer.get_by_role("button", name="Seal key").click()
    page.keyboard.press("Escape")


# -- every tool, no key ----------------------------------------------------------------

@pytest.mark.parametrize("slug", BENTO_TOOLS)
def test_tool_works_with_the_mousy_sample_and_no_key(ui: Page, site_url: str, slug: str) -> None:
    errors: list[str] = []
    ui.on("console", lambda msg: errors.append(msg.text) if msg.type == "error" else None)
    ui.on("pageerror", lambda err: errors.append(str(err)))
    load_sample(ui, site_url, slug)

    expect(ui.get_by_test_id("answer")).not_to_have_text("–")
    expect(ui.get_by_test_id("run-log")).to_contain_text(re.compile(r"0 AI calls · \d+ ms"))
    expect(ui.get_by_test_id("ai-label")).to_have_count(0)            # nothing claims to be AI
    expect(ui.locator(".bento__exports button").first).to_be_visible()
    assert errors == []


# -- Bug Finder ---------------------------------------------------------------------------

def test_bug_finder_matches_the_answer_key_and_stamps_fails(ui: Page, site_url: str) -> None:
    load_sample(ui, site_url, "bug-finder")
    expect(ui.get_by_test_id("key-check")).to_contain_text("Matches answer key")
    findings = ui.get_by_test_id("finding")
    expect(findings).to_have_count(2)
    expect(findings.first).to_contain_text("✕Fail")                 # icon and word, not colour only
    expect(ui.locator(".hanko")).to_have_count(2)
    assert ui.locator(".hanko").first.evaluate("el => getComputedStyle(el).animationName") == "slam"


def test_hanko_does_not_slam_with_reduced_motion(browser: Browser, site_url: str) -> None:
    context = browser.new_context(reduced_motion="reduce")
    page = context.new_page()
    load_sample(page, site_url, "bug-finder")
    expect(page.locator(".hanko").first).to_be_visible()
    assert page.locator(".hanko").first.evaluate("el => getComputedStyle(el).animationName") == "none"
    context.close()


def test_bug_finder_bulk_csv_matches_the_whole_answer_key(ui: Page, site_url: str) -> None:
    load_sample(ui, site_url, "bug-finder", "mousy_conversations.csv")
    expect(ui.get_by_test_id("bulk-findings")).to_be_visible()
    total = len(ANSWER_KEY["conversations"])
    expect(ui.get_by_test_id("results")).to_contain_text(f"{total} of {total} conversations match the Mousy answer key exactly.")


def test_bug_reports_are_written_as_steps(ui: Page, site_url: str) -> None:
    load_sample(ui, site_url, "bug-finder")
    name, data = download(ui, "Download bug reports (.md)")
    text = data.decode("utf-8")
    assert name == "01_cancel_order_chat_bug_reports.md"
    for heading in ("### Steps", "### Expected", "### Actual", "### Evidence"):
        assert heading in text


# -- Transcript Converter --------------------------------------------------------------

def test_converter_yaml_round_trips(ui: Page, site_url: str) -> None:
    load_sample(ui, site_url, "transcript-converter")
    steps = ui.get_by_test_id("answer").inner_text()
    ui.get_by_test_id("out-format").select_option("yaml")
    ui.get_by_test_id("run").click()
    expect(ui.get_by_test_id("script-preview")).to_contain_text('format: "kitsune-test-script/1"')
    name, data = download(ui, "Download .yaml")
    assert name.endswith(".yaml")

    # Feed the YAML back in: the same steps come out.
    ui.get_by_role("button", name="Clear").click()
    ui.get_by_label("Conversation").fill(data.decode("utf-8"))
    ui.get_by_test_id("out-format").select_option("json")
    ui.get_by_test_id("run").click()
    expect(ui.get_by_test_id("answer")).to_have_text(steps)
    script = json.loads(ui.get_by_test_id("script-preview").inner_text())
    assert script["format"] == "kitsune-test-script/1"


def test_converter_order_ids_become_a_variable(ui: Page, site_url: str) -> None:
    ui.goto(f"{site_url}tools/transcript-converter/")
    ui.get_by_label("Conversation").fill("Mousy Assist: Which order?\nCustomer: it's MO-4417")
    ui.get_by_test_id("run").click()
    expect(ui.get_by_test_id("script-preview")).to_contain_text("{{order_id}}")
    ui.get_by_test_id("order-variable").uncheck()
    ui.get_by_test_id("run").click()
    expect(ui.get_by_test_id("script-preview")).to_contain_text("MO-4417")


def test_converter_reads_the_sample_zip(ui: Page, site_url: str) -> None:
    load_sample(ui, site_url, "transcript-converter", "mousy_sample_pack.zip")
    expect(ui.get_by_test_id("loaded-file")).to_be_visible()
    files = ui.get_by_test_id("converted-files").locator("tbody tr td:first-child").all_inner_texts()
    assert len(files) >= 8
    assert len(files) == len(set(files)), "two scripts would overwrite each other in the zip"
    name, data = download(ui, re.compile(r"^Download \.zip \(\d+ \.txt files\)$"))
    assert name.endswith(".zip")
    assert len(zipfile.ZipFile(io.BytesIO(data)).namelist()) == len(files)


# -- First Utterance Finder and Metrics Extractor --------------------------------------

def test_first_utterances_match_the_answer_key(ui: Page, site_url: str) -> None:
    load_sample(ui, site_url, "first-utterance-finder")
    total = len(ANSWER_KEY["conversations"])
    expect(ui.get_by_test_id("key-check")).to_contain_text(f"{total} of {total} match the Mousy answer key.")
    name, data = download(ui, "Download .xlsx")
    sheet = openpyxl.load_workbook(io.BytesIO(data)).active
    assert sheet.max_row == total + 1


def test_metrics_cross_check_and_csv_columns(ui: Page, site_url: str) -> None:
    load_sample(ui, site_url, "metrics-extractor")
    expect(ui.get_by_test_id("cross-check")).to_contain_text("✓Agrees")
    _, data = download(ui, "Download .csv")
    header = data.decode("utf-8").splitlines()[0].split(",")
    assert "metric_fallbacks" in header and "metric_ending" in header


# -- Test Case Generator -----------------------------------------------------------------

def test_test_cases_use_the_column_mapping(ui: Page, site_url: str) -> None:
    load_sample(ui, site_url, "test-case-generator")
    expect(ui.get_by_test_id("map-transcript")).to_have_value("transcript")
    expect(ui.get_by_test_id("map-id")).to_have_value("conversation_id")
    expect(ui.get_by_test_id("answer")).to_have_text(str(len(ANSWER_KEY["conversations"])))

    # Map the ID to another column: the case IDs follow it.
    ui.get_by_test_id("map-id").select_option("started_at")
    ui.get_by_test_id("run").click()
    expect(ui.get_by_test_id("test-cases").locator("tbody tr").first).not_to_contain_text("CNV-")

    _, data = download(ui, "Download .xlsx")
    sheet = openpyxl.load_workbook(io.BytesIO(data)).active
    header = [c.value for c in sheet[1]]
    assert {"test_case_id", "steps", "expected_results", "priority"} <= set(header)


# -- Utterance Generator -----------------------------------------------------------------

def test_utterances_from_the_intents_workbook(ui: Page, site_url: str) -> None:
    load_sample(ui, site_url, "utterance-generator")
    options = ui.get_by_test_id("intent-pick").locator("option")
    expect(options).to_have_count(11)                                   # "All" + 10 intents
    first = ui.get_by_test_id("utterances").locator("tbody").inner_text()

    # Same input, same output: the rules are seeded, so tests and visitors see stable results.
    ui.get_by_test_id("run").click()
    expect(ui.get_by_test_id("run-status")).to_have_text("Done")
    assert ui.get_by_test_id("utterances").locator("tbody").inner_text() == first

    ui.get_by_test_id("intent-pick").select_option(index=1)
    ui.get_by_test_id("count-pick").select_option("10")
    ui.get_by_test_id("run").click()
    expect(ui.get_by_test_id("answer")).to_have_text("10")


# -- Report Builder and Email Templates ------------------------------------------------

def test_report_ranks_by_impact_and_exports_safe_html(ui: Page, site_url: str) -> None:
    load_sample(ui, site_url, "report-builder")
    clean = sum(1 for c in ANSWER_KEY["conversations"] if not c["bugs"])
    expect(ui.get_by_test_id("answer")).to_have_text(f"{round(100 * clean / len(ANSWER_KEY['conversations']))}%")
    expect(ui.get_by_test_id("report")).to_contain_text("Top issues (value over volume)")
    _, data = download(ui, "Download .html")
    html = data.decode("utf-8")
    assert "<script" not in html.lower()
    assert "All Mousy data is invented" in html


def test_critical_email_names_the_worst_issue(ui: Page, site_url: str) -> None:
    load_sample(ui, site_url, "email-templates")
    ui.get_by_test_id("template").select_option("critical")
    ui.get_by_test_id("run").click()
    expect(ui.get_by_test_id("email-subject")).to_contain_text("Subject: Critical: Fallback loop")
    _, data = download(ui, "Download .eml")
    eml = data.decode("utf-8")
    assert "Subject: Critical:" in eml and "To: undisclosed-recipients:;" in eml


def test_custom_email_asks_for_a_key(ui: Page, site_url: str) -> None:
    load_sample(ui, site_url, "email-templates")
    ui.get_by_test_id("template").select_option("custom")
    ui.get_by_test_id("run").click()
    expect(ui.get_by_test_id("needs-key")).to_contain_text("Needs AI key")


# -- QA Handbook ----------------------------------------------------------------------------

def test_handbook_checklist_counts_and_persists(ui: Page, site_url: str) -> None:
    ui.goto(f"{site_url}tools/qa-handbook/")
    expect(ui.get_by_test_id("checklist-count")).to_have_text("0 of 9 steps done")
    ui.locator("#hb-checklist input[type=checkbox]").nth(0).check()
    ui.locator("#hb-checklist input[type=checkbox]").nth(1).check()
    expect(ui.get_by_test_id("checklist-count")).to_have_text("2 of 9 steps done")
    ui.reload()
    expect(ui.get_by_test_id("checklist-count")).to_have_text("2 of 9 steps done")


def test_handbook_sample_maps_findings_to_practices(ui: Page, site_url: str) -> None:
    ui.goto(f"{site_url}tools/qa-handbook/")
    ui.get_by_test_id("load-sample").click()
    expect(ui.get_by_test_id("results")).to_contain_text("transcripts/06")
    expect(ui.get_by_test_id("finding").first).to_be_visible()


# -- the AI path (mocked provider) -----------------------------------------------------

def gemini_reply(payload: dict) -> str:
    return json.dumps({"candidates": [{"content": {"parts": [{"text": json.dumps(payload)}]}}]})


def test_ai_review_sends_the_key_only_to_the_provider(ui: Page, site_url: str) -> None:
    site_host = urlparse(site_url).netloc
    calls: list[dict] = []
    seen: list[str] = []

    def provider(route: Route) -> None:
        request = route.request
        calls.append({"url": request.url, "headers": request.headers, "body": request.post_data or ""})
        route.fulfill(status=200, content_type="application/json", body=gemini_reply({
            "summary": "The bot misread a cancel request as a refund.",
            "sentiment": {"start": "neutral", "end": "neutral"},
            "findings": [
                {"type": "misroute", "turn": 3, "severity": "high", "title": "Wrong intent", "evidence": "x", "fix": "y"},
                {"type": "tone", "turn": 13, "severity": "low", "title": "Abrupt ending", "evidence": "z", "fix": ""},
            ],
        }))

    ui.route(f"https://{GEMINI_HOST}/**", provider)
    ui.on("request", lambda req: seen.append(req.url))
    ui.on("console", lambda msg: seen.append(msg.text))
    ui.goto(f"{site_url}tools/bug-finder/")
    add_key(ui)
    expect(ui.get_by_test_id("use-ai")).to_be_checked()
    ui.get_by_test_id("load-sample").click()
    expect(ui.get_by_test_id("run-status")).to_have_text("Done")

    expect(ui.get_by_test_id("ai-summary")).to_contain_text("misread a cancel request")
    expect(ui.get_by_test_id("ai-label").first).to_contain_text("AI")
    expect(ui.get_by_text("AI agrees")).to_be_visible()                  # rules + AI on the same bug
    expect(ui.get_by_test_id("finding")).to_have_count(3)                 # 2 rules + 1 new from AI
    expect(ui.get_by_test_id("run-log")).to_contain_text("1 AI call ·")

    assert len(calls) == 1
    assert calls[0]["headers"].get("x-goog-api-key") == FAKE_KEY          # the key goes in the header...
    assert FAKE_KEY not in calls[0]["url"] and FAKE_KEY not in calls[0]["body"]   # ...and nowhere else
    others = [u for u in seen if u.startswith("http") and urlparse(u).netloc not in (site_host, GEMINI_HOST)]
    assert others == []
    assert not any(FAKE_KEY in line for line in seen)


def test_ai_failure_falls_back_to_rules(ui: Page, site_url: str) -> None:
    ui.route(f"https://{GEMINI_HOST}/**", lambda route: route.fulfill(status=401, body='{"error": {"message": "API key not valid"}}'))
    ui.goto(f"{site_url}tools/bug-finder/")
    add_key(ui)
    ui.get_by_test_id("load-sample").click()
    expect(ui.get_by_test_id("run-status")).to_have_text("Done")
    expect(ui.get_by_test_id("answer")).to_have_text("2")
    expect(ui.get_by_test_id("run-log")).to_contain_text("skipped")
    assert FAKE_KEY not in ui.get_by_test_id("run-log").inner_text()


def test_ai_report_summary_sends_totals_not_transcripts(ui: Page, site_url: str) -> None:
    bodies: list[str] = []

    def provider(route: Route) -> None:
        bodies.append(route.request.post_data or "")
        route.fulfill(status=200, content_type="application/json", body=gemini_reply({"summary": "Mostly healthy; fix wrong intents first."}))

    ui.route(f"https://{GEMINI_HOST}/**", provider)
    ui.goto(f"{site_url}tools/report-builder/")
    add_key(ui)
    ui.get_by_test_id("load-sample").click()
    expect(ui.get_by_test_id("report")).to_contain_text("Mostly healthy")
    expect(ui.get_by_test_id("ai-label").first).to_be_visible()
    assert len(bodies) == 1
    first_evidence = next(b["evidence"] for c in ANSWER_KEY["conversations"] for b in c["bugs"])
    assert first_evidence not in bodies[0]
