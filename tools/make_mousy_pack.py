"""Build the Mousy sample pack into app/public/sample-data/. Everything is invented.

This script never reads a real export. The few shape numbers it needs (channel mix, turn
lengths, timeout length) were copied by hand from the private profile as the constants below.
The seed is fixed and every file is written byte-for-byte reproducibly, so a re-run only
changes files when this script changes. Scenario list: sample-data/SCENARIOS.md.

Run:  python tools/make_mousy_pack.py            (writes the default folder)
      python tools/make_mousy_pack.py --out DIR  (writes somewhere else, used by the tests)
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import random
import re
import shutil
import zipfile
from dataclasses import dataclass, field
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_OUT = ROOT / "app" / "public" / "sample-data"

SEED = 4417
BOT = "Mousy Assist"
CUSTOMER = "Customer"

# Shape constants, copied from the private profile (numbers only).
TOTAL_VARIATIONS = 32
VARIATION_CHAT_COUNT = 4          # with the 3 core chats: 7 of 40, close to the profile's 16%
CHAT_TIMEOUT_SEC = 600
VOICE_SILENCE_SEC = 8
BOT_WORDS_PER_SEC = 2.6
CUSTOMER_WORDS_PER_SEC = 2.2
VOICE_REPROMPT_CHANCE = 0.3
VOICE_FILLER_OPEN_CHANCE = 0.2
WINDOW_START = datetime(2025, 3, 1, tzinfo=timezone.utc)
WINDOW_DAYS = 14

# Fixed timestamp for files inside .xlsx and .zip archives, so builds are reproducible.
ARCHIVE_TIME = (2025, 3, 14, 0, 0, 0)

FIRST_NAMES = ["Pip", "Juniper", "Wren", "Clover", "Hazel", "Basil", "Poppy", "Fennel",
               "Sorrel", "Bramble", "Tansy", "Moss"]
ITEMS = [("blue kettle", 34.00), ("two-slice toaster", 29.50), ("bamboo cutting board", 18.00),
         ("milk frother", 22.00), ("cheese grater", 9.50), ("desk lamp", 41.00),
         ("set of mixing bowls", 26.00), ("phone stand", 12.00), ("night light", 15.00),
         ("egg timer", 7.50)]
STREETS = ["Cheddar Row", "Whisker Lane", "Pantry Court", "Crumb Street", "Burrow Close", "Acorn Way"]
TOWNS = ["Nibbleton", "Burrowford", "Crumbridge", "Hollowmere"]
CARRIER = "Pawcel Post"
DIGIT_WORDS = ["zero", "one", "two", "three", "four", "five", "six", "seven", "eight", "nine"]

INTENTS = {
    "cancel_order": "Cancel an order that has not shipped yet.",
    "refund_status": "Find out where a refund is and when it will arrive.",
    "change_delivery_address": "Send an order that has not arrived yet to a different address.",
    "billing_dispute": "Question a charge, for example being charged twice for one order.",
    "track_parcel": "Find out where a parcel is and when it will be delivered.",
    "gift_card_balance": "Check how much money is left on a Mousy gift card.",
    "return_damaged_item": "Return an item that arrived broken and get a replacement or refund.",
    "talk_to_human": "Speak to a person instead of the bot.",
    "update_payment_method": "Change the payment card saved on the account.",
    "product_question": "Ask about a product, such as care instructions or warranty.",
}

# Customer opening requests. Voice callers are short; chat users write a little more.
REQUESTS = {
    "cancel_order": {
        "voice": ["cancel my order", "I want to cancel an order", "cancel order"],
        "chat": ["Hi, I want to cancel my order", "I need to cancel an order I placed yesterday",
                 "can I cancel my order please"]},
    "refund_status": {
        "voice": ["where's my refund", "refund status", "I'm calling about a refund"],
        "chat": ["I returned something two weeks ago and still have no refund", "Where is my refund?",
                 "checking on a refund"]},
    "change_delivery_address": {
        "voice": ["change my delivery address", "I need to change the address", "wrong address on my order"],
        "chat": ["I put the wrong address on my order", "Can I change the delivery address?",
                 "need to update where my order is going"]},
    "billing_dispute": {
        "voice": ["I was charged twice", "there's a wrong charge on my card", "billing problem"],
        "chat": ["I was charged twice for one order", "There's a charge on my card I don't recognise",
                 "my card got billed two times"]},
    "track_parcel": {
        "voice": ["where's my package", "track my order", "my parcel hasn't come"],
        "chat": ["Where is my parcel?", "can you track my order", "my order still hasn't arrived"]},
    "gift_card_balance": {
        "voice": ["gift card balance", "how much is on my gift card", "check a gift card"],
        "chat": ["How much is left on my gift card?", "gift card balance please", "can I check my gift card"]},
    "return_damaged_item": {
        "voice": ["my order came broken", "return a damaged item", "something arrived damaged"],
        "chat": ["One of my items arrived broken", "I want to return a damaged item",
                 "the box was crushed and the item is broken"]},
    "talk_to_human": {
        "voice": ["agent", "representative", "talk to a person"],
        "chat": ["can I talk to a real person", "human please", "I want to speak to an agent"]},
    "update_payment_method": {
        "voice": ["update my card", "change payment method", "new card on my account"],
        "chat": ["I need to update my payment card", "How do I change my payment method?",
                 "my card expired, can I add a new one"]},
    "product_question": {
        "voice": ["is the {item} dishwasher safe", "question about a product",
                  "does the {item} come with a warranty"],
        "chat": ["Is the {item} dishwasher safe?", "Does the {item} have a warranty?",
                 "what's the {item} made of"]},
}

# Wrong route the bot takes for a misroute, and how the customer pushes back.
MISROUTES = {
    "cancel_order": ("refund_status", "I can help with refunds! Which order should I refund?",
                     "no, cancel. not a refund"),
    "refund_status": ("cancel_order", "I can help you cancel an order. Which order would you like to cancel?",
                      "no, I don't want to cancel, I want my refund"),
    "change_delivery_address": ("track_parcel", "Sure, let me track your parcel. What's the order number?",
                                "no, I need to change the address, not track it"),
    "return_damaged_item": ("cancel_order", "I can help cancel an order. Which order is it?",
                            "no, it already arrived. it's broken"),
    "track_parcel": ("change_delivery_address", "Sure. What's the new delivery address?",
                     "no, I just want to know where it is"),
}

FILLER_OPENERS = ["hello?", "um, hi", "yes hi"]
YES = ["yes", "yeah", "correct", "yes please", "yep"]
THANKS_BYE = ["no, that's all", "no thanks", "nope, that's it. thanks"]

BUG_INFO = {
    "misroute": ("high", "The bot sent the customer down the wrong intent."),
    "fallback_loop": ("high", "The bot re-prompted 3 or more times in a row without progress."),
    "premature_escalation": ("medium", "The bot transferred to a person before trying the self-service flow."),
    "timeout_no_goodbye": ("low", "The customer went quiet and the conversation closed with no closing message."),
    "bot_repeats": ("low", "The bot sent the same information twice."),
    "wrong_order_id": ("high", "The bot read back or used a different order ID than the customer gave."),
    "ignored_correction": ("high", "The customer corrected the bot and the bot carried on with the old value."),
}

# Which intents each bug can sensibly appear in.
BUG_INTENTS = {
    "misroute": list(MISROUTES),
    "fallback_loop": [i for i in INTENTS if i != "talk_to_human"],
    "premature_escalation": ["cancel_order", "billing_dispute", "return_damaged_item", "change_delivery_address",
                             "refund_status"],
    "timeout_no_goodbye": list(INTENTS),
    "bot_repeats": ["refund_status", "track_parcel", "gift_card_balance", "product_question"],
    "wrong_order_id": ["cancel_order", "refund_status", "change_delivery_address", "track_parcel",
                       "return_damaged_item"],
    "ignored_correction": ["change_delivery_address", "return_damaged_item"],
}

VARIATION_BUG_BAG = (["misroute"] * 4 + ["fallback_loop"] * 2 + ["premature_escalation"] * 3
                     + ["timeout_no_goodbye"] * 2 + ["bot_repeats"] * 3 + ["wrong_order_id"] * 2
                     + ["ignored_correction"] * 3 + [None] * 13)

QUEUE_FOR_INTENT = {"billing_dispute": "billing_team", "return_damaged_item": "returns_desk"}


@dataclass
class Spec:
    number: int
    intent: str
    channel: str
    bugs: list[str]
    ending: str = "resolved"           # resolved | timeout | hangup
    core_file: str | None = None

    @property
    def conversation_id(self) -> str:
        return f"CNV-{self.number:06d}"


CORE_SPECS = [
    Spec(1, "cancel_order", "chat", ["misroute", "ignored_correction"], core_file="01_cancel_order_chat"),
    Spec(2, "refund_status", "voice", [], core_file="02_refund_status_voice"),
    Spec(3, "change_delivery_address", "voice", ["wrong_order_id"], core_file="03_change_address_voice"),
    Spec(4, "billing_dispute", "voice", ["premature_escalation"], core_file="04_billing_dispute_voice"),
    Spec(5, "track_parcel", "voice", ["fallback_loop"], core_file="05_track_parcel_voice"),
    Spec(6, "gift_card_balance", "voice", ["bot_repeats"], core_file="06_gift_card_voice"),
    Spec(7, "return_damaged_item", "chat", ["ignored_correction"], core_file="07_return_damaged_chat"),
    Spec(8, "talk_to_human", "chat", ["timeout_no_goodbye"], ending="timeout", core_file="08_talk_to_human_chat"),
]


@dataclass
class Facts:
    """Everything invented about one customer and order. Drawn once per conversation."""
    name: str
    order_digits: str
    phone: str
    item: str
    item_price: float
    item2: str
    address: str
    address2: str
    refund_day: str
    ship_day: str
    arrive_weekday: str
    town: str
    card_last4: str
    card_balance: float
    wait_minutes: int
    lookup_ms: int
    confidence: float

    @property
    def order_id(self) -> str:
        return f"MO-{self.order_digits}"


@dataclass
class Turn:
    speaker: str
    text: str
    tags: list[str] = field(default_factory=list)
    wait_before: float | None = None   # forced pause (silence before a timeout prompt)


class Conversation:
    def __init__(self, spec: Spec, facts: Facts, rng: random.Random, bugs: list[str]):
        self.spec = spec
        self.f = facts
        self.rng = rng
        self.bugs = set(bugs)
        self.voice = spec.channel == "voice"
        self.turns: list[Turn] = []
        self.steps: list[tuple[str, int]] = []
        self.used_order_id = facts.order_id
        self.escalation: tuple[str, str] | None = None   # (queue, reason)
        self.fallbacks = 0
        self.auth = "skipped"

    # -- primitives -------------------------------------------------------------
    def bot(self, text: str, *tags: str, wait_before: float | None = None) -> None:
        self.turns.append(Turn(BOT, text, list(tags), wait_before))

    def cust(self, text: str, *tags: str) -> None:
        self.turns.append(Turn(CUSTOMER, text, list(tags)))

    def step(self, name: str) -> None:
        self.steps.append((name, len(self.turns)))

    def pick(self, options: list[str]) -> str:
        return self.rng.choice(options)

    @property
    def done(self) -> bool:
        return self.escalation is not None

    # -- shared building blocks -------------------------------------------------
    def request_text(self, intent: str, avoid: str | None = None) -> str:
        options = [o.format(item=self.f.item) for o in REQUESTS[intent][self.spec.channel]]
        if avoid in options and len(options) > 1:
            options.remove(avoid)
        return self.pick(options)

    def open(self) -> None:
        self.step("greet")
        if self.voice:
            self.bot("Thanks for calling Mousy. You're speaking with Mousy Assist, a virtual assistant. "
                     "In a few words, how can I help today?")
            if self.rng.random() < VOICE_FILLER_OPEN_CHANCE:
                self.cust(self.pick(FILLER_OPENERS))
                self.bot("Hi there. What can I help you with today?")
        else:
            self.bot("Hi, I'm Mousy Assist. How can I help with your order today?")
        self.cust(self.request_text(self.spec.intent), "first_utterance")

    def fallback_loop(self) -> None:
        self.step("fallback")
        said = self.turns[-1].text
        for attempt in range(3):
            tags = ("fallback_loop",) if attempt == 2 else ()
            self.bot("Sorry, I didn't catch that. Could you say it another way?", *tags)
            self.fallbacks += 1
            said = self.request_text(self.spec.intent, avoid=said)
            self.cust(said)
        self.escalate("general_support", "repeated_fallback",
                      "I'm having trouble understanding. Let me connect you with someone who can help.")

    def premature_escalation(self) -> None:
        queue = QUEUE_FOR_INTENT.get(self.spec.intent, "general_support")
        self.escalate(queue, "bot_transfer",
                      "Let me transfer you to a specialist who can help with that. Please hold.",
                      "premature_escalation")

    def escalate(self, queue: str, reason: str, message: str, *tags: str) -> None:
        self.step("escalate")
        self.bot(message, *tags)
        self.escalation = (queue, reason)

    def misroute(self) -> None:
        routed, bot_line, pushback = MISROUTES[self.spec.intent]
        self.step(routed)
        self.bot(bot_line, "misroute")
        self.cust(pushback if self.voice else pushback.replace("no, cancel.", "no, CANCEL."))
        if "ignored_correction" in self.bugs and self.spec.intent not in BUG_INTENTS["ignored_correction"]:
            self.bot("Sorry, I didn't get that. Could you say it another way?", "ignored_correction")
            self.fallbacks += 1
            first = next(t.text for t in self.turns if "first_utterance" in t.tags)
            self.cust(self.request_text(self.spec.intent, avoid=first))
        else:
            self.bot("Sorry about that. Let's sort that out instead.")

    def spoken(self, digits: str) -> str:
        return " ".join(DIGIT_WORDS[int(d)] for d in digits)

    def identify_order(self) -> None:
        self.step("identify_order")
        digits = self.f.order_digits
        wrong = "wrong_order_id" in self.bugs
        if wrong:
            swapped = digits[:2] + digits[3] + digits[2]
            bad = swapped if swapped != digits else digits[:3] + str((int(digits[3]) + 3) % 10)
            self.used_order_id = f"MO-{bad}"
        if self.voice:
            self.bot("Sure. What's the order number? It starts with M O.")
            if self.rng.random() < 0.25:
                self.cust("hold on, let me find it")
                self.bot("No problem. Take your time.")
            self.cust(f"M O {self.spoken(digits)}")
            if self.rng.random() < VOICE_REPROMPT_CHANCE:
                self.bot("Sorry, I didn't catch that. Please say the order number again.")
                self.cust(f"M O {self.spoken(digits)}")
            heard = self.used_order_id.removeprefix("MO-")
            self.bot(f"Thanks. I heard M O {self.spoken(heard)}. Is that right?",
                     *(("wrong_order_id",) if wrong else ()))
            self.cust(self.pick(YES))
        else:
            self.bot("Sure. What's your order number? It starts with MO.")
            self.cust(self.pick([self.f.order_id, f"it's {self.f.order_id}", self.f.order_id.lower()]))
            self.bot(f"Thanks. Looking up order {self.used_order_id} now.",
                     *(("wrong_order_id",) if wrong else ()))

    def verify_caller(self) -> None:
        if not self.voice:
            return
        self.step("verify_caller")
        phone_digits = self.f.phone.replace("-", "")
        self.bot("For security, please say the phone number on the account.")
        self.cust(f"{self.spoken(phone_digits[:3])}, {self.spoken(phone_digits[3:])}")
        self.bot(f"Thanks. I have {self.f.phone}. One moment while I find your details.")
        if self.rng.random() < 0.5:
            self.bot("And to confirm, what's the first name on the account?")
            self.cust(self.f.name.lower())
        self.bot(f"Thanks, {self.f.name}. I found your account.")
        self.auth = "passed"

    def info(self, text: str) -> None:
        """The key answer of a flow. With the bot_repeats bug, the bot sends it twice."""
        self.bot(text)
        if "bot_repeats" in self.bugs:
            self.cust(self.pick(["okay", "okay, thanks", "right"]))
            self.bot(text, "bot_repeats")
            self.cust("yes, you said that")
        if self.voice and self.rng.random() < 0.5:
            self.bot("Would you like me to text you a summary?")
            if self.rng.random() < 0.7:
                self.cust(self.pick(YES))
                self.bot(f"Done. I've sent it to {self.f.phone}.")
            else:
                self.cust("no, that's fine")

    def wrap_up(self) -> None:
        self.step("offer_more_help")
        self.bot("Is there anything else I can help with today?")
        self.cust(self.pick(THANKS_BYE))
        if self.voice and self.rng.random() < 0.6:
            self.step("rate_call")
            self.bot("Before you go, could you rate this call from one to five?")
            self.cust(self.pick(["five", "four", "three", "five, thanks"]))
            self.bot("Thank you for the feedback.")
        self.step("close")
        if self.voice:
            self.bot("Thanks for calling Mousy. Have a lovely day. Goodbye.")
        else:
            self.bot("Thanks for chatting with Mousy. Have a lovely day!")

    # -- intent flows -----------------------------------------------------------
    def flow_cancel_order(self) -> None:
        self.identify_order()
        self.verify_caller()
        self.step("lookup_order")
        self.bot(f"Order {self.used_order_id} has the {self.f.item} and hasn't shipped yet. "
                 "Do you want to cancel the whole order?")
        self.cust(self.pick(YES))
        self.step("cancel_order")
        self.info(f"Done. Order {self.used_order_id} is cancelled. ${self.f.item_price:.2f} goes back to your "
                  "card in 3 to 5 business days.")

    def flow_refund_status(self) -> None:
        self.identify_order()
        self.verify_caller()
        self.step("lookup_refund")
        self.info(f"I see a refund of ${self.f.item_price:.2f} for order {self.used_order_id}, issued on "
                  f"{self.f.refund_day}. It can take up to 5 business days to show on your statement.")
        self.cust(self.pick(["okay so it's coming?", "so when will I see it", "okay"]))
        self.bot("Yes, it's on its way. You'll get an email when your bank confirms it.")

    def flow_change_delivery_address(self) -> None:
        self.identify_order()
        self.verify_caller()
        self.step("lookup_order")
        self.bot(f"Order {self.used_order_id} is due on {self.f.arrive_weekday}. What's the new delivery address?")
        self.cust(self.f.address)
        self.bot(f"Got it: {self.f.address}. Should I update the order?")
        if "ignored_correction" in self.bugs:
            self.cust(f"wait, no, sorry. it's {self.f.address2}")
            self.step("update_address")
            self.bot(f"Done. Order {self.used_order_id} will go to {self.f.address} on {self.f.arrive_weekday}.",
                     "ignored_correction")
            return
        self.cust(self.pick(YES))
        self.step("update_address")
        self.info(f"Done. Order {self.used_order_id} will go to {self.f.address} on {self.f.arrive_weekday}.")

    def flow_billing_dispute(self) -> None:
        self.identify_order()
        self.verify_caller()
        self.step("check_charges")
        self.info(f"I see two charges of ${self.f.item_price:.2f} for order {self.used_order_id}. One is a "
                  "temporary hold. It drops off within 2 business days.")
        if self.rng.random() < 0.5:
            self.cust("okay but I want someone to check it")
            self.escalate("billing_team", "billing_review",
                          "Of course. I'm passing this to our billing team with the details. Connecting you now.")
        else:
            self.cust("ah okay, that makes sense")

    def flow_track_parcel(self) -> None:
        self.identify_order()
        self.step("track_parcel")
        self.info(f"Order {self.used_order_id} shipped on {self.f.ship_day} with {CARRIER}. It's in {self.f.town} "
                  f"and should arrive on {self.f.arrive_weekday} by 6 pm.")
        if self.rng.random() < 0.5:
            self.cust("can you send me the tracking link")
            self.bot("Sure. I've sent the tracking link to the email on your account.")
        else:
            self.cust(self.pick(["great", "okay thanks", "perfect"]))

    def flow_gift_card_balance(self) -> None:
        self.step("check_gift_card")
        verb = "say" if self.voice else "type"
        self.bot(f"Sure. Please {verb} the last four digits of the gift card.")
        self.cust(self.spoken(self.f.card_last4) if self.voice else self.f.card_last4)
        self.info(f"Your gift card ending in {self.f.card_last4} has ${self.f.card_balance:.2f} left. "
                  "It doesn't expire.")

    def flow_return_damaged_item(self) -> None:
        self.identify_order()
        self.verify_caller()
        self.step("lookup_order")
        self.bot(f"I'm sorry it arrived damaged. Order {self.used_order_id} has the {self.f.item} and the "
                 f"{self.f.item2}. Which one is damaged?")
        returned = self.f.item
        if "ignored_correction" in self.bugs:
            self.cust(f"the {self.f.item2}")
            self.bot(f"The {self.f.item2}, got it.")
            self.cust(f"sorry, no. I mean the {self.f.item}")
            returned = self.f.item2
            tags: tuple[str, ...] = ("ignored_correction",)
        else:
            self.cust(f"the {self.f.item}")
            tags = ()
        self.step("start_return")
        self.bot(f"I've started a return for the {returned}. A prepaid label is on its way to your email. "
                 "Would you like a replacement or a refund?", *tags)
        choice = self.pick(["replacement", "refund", "a replacement please"])
        self.cust(choice)
        if "refund" in choice:
            self.info(f"Okay. You'll get ${self.f.item_price:.2f} back once the courier scans the parcel.")
        else:
            self.info("Okay. A new one ships as soon as the courier scans the return parcel.")

    def flow_talk_to_human(self) -> None:
        self.step("offer_self_service")
        self.bot("I can help with most order questions. What's it about?")
        self.cust(self.pick(["just get me a person please", "I want a human", "person please"]))
        self.escalate("general_support", "customer_request",
                      f"Sure. I'm connecting you to a Mousy support specialist. The wait is about "
                      f"{self.f.wait_minutes} minutes.")

    def flow_update_payment_method(self) -> None:
        self.verify_caller()
        self.step("send_secure_link")
        self.info("For your security, I can't take card details here. I've sent a secure link to the email "
                  "on your account. It works for 24 hours.")
        self.cust(self.pick(["okay", "got it, thanks", "okay I see it"]))

    def flow_product_question(self) -> None:
        self.step("product_lookup")
        self.info(f"The {self.f.item} has a one-year Mousy warranty. It isn't dishwasher safe, so please "
                  "hand-wash it in warm soapy water.")

    # -- assembly ---------------------------------------------------------------
    def build(self) -> None:
        self.open()
        if "premature_escalation" in self.bugs:
            self.premature_escalation()
            return
        if "fallback_loop" in self.bugs:
            self.fallback_loop()
            return
        if "misroute" in self.bugs:
            self.misroute()
        getattr(self, f"flow_{self.spec.intent}")()
        if not self.done:
            self.wrap_up()

    def apply_ending(self, ending: str) -> None:
        """Cut a conversation short for a timeout or hang-up. The cut lands on a bot question."""
        if ending == "resolved":
            return
        questions = [i for i, t in enumerate(self.turns) if t.speaker == BOT and t.text.endswith("?") and i >= 1]
        later = [i for i in questions if i > len(self.turns) // 3] or questions
        cut = later[0] if ending == "timeout" else self.rng.choice(later)
        # Everything after the cut never happened, including a later escalation.
        self.turns = self.turns[: cut + 1]
        self.steps = [(name, at) for name, at in self.steps if at <= cut]
        if not any(name == "escalate" for name, _ in self.steps):
            self.escalation = None
        if ending == "hangup":
            return
        silence = CHAT_TIMEOUT_SEC // 5 if not self.voice else VOICE_SILENCE_SEC
        no_goodbye = "timeout_no_goodbye" in self.bugs
        self.bot("Are you still there?", *(("timeout_no_goodbye",) if no_goodbye else ()), wait_before=silence)
        if not no_goodbye:
            self.step("close")
            what, again = ("call", "call back") if self.voice else ("chat", "start a new chat")
            self.bot(f"I haven't heard back, so I'll close this {what} now. You're welcome to {again} any time. "
                     "Goodbye.", wait_before=silence)


def draw_facts(rng: random.Random, start: datetime) -> Facts:
    item, price = rng.choice(ITEMS)
    item2 = rng.choice([i for i, _ in ITEMS if i != item])
    street, street2 = rng.sample(STREETS, 2)
    town = rng.choice(TOWNS)
    house = rng.randint(2, 98)
    refund = start - timedelta(days=rng.randint(2, 6))
    ship = start - timedelta(days=rng.randint(1, 3))
    arrive = start + timedelta(days=rng.randint(1, 4))
    return Facts(
        name=rng.choice(FIRST_NAMES),
        order_digits=f"{rng.randint(1000, 9999)}",
        phone=f"555-01{rng.randint(0, 99):02d}",
        item=item,
        item_price=price,
        item2=item2,
        address=f"{house} {street}, {town}",
        address2=f"{house + rng.randint(1, 9)} {street2}, {town}",
        refund_day=f"{refund:%B} {refund.day}",
        ship_day=f"{ship:%B} {ship.day}",
        arrive_weekday=f"{arrive:%A}",
        town=rng.choice(TOWNS),
        card_last4=f"{rng.randint(0, 9999):04d}",
        card_balance=round(rng.randint(5, 150) + rng.choice([0, 0.25, 0.5, 0.75]), 2),
        wait_minutes=rng.randint(2, 9),
        lookup_ms=rng.randint(180, 1400),
        confidence=round(rng.uniform(0.62, 0.98), 2),
    )


def turn_times(conv: Conversation, rng: random.Random) -> tuple[list[int], int]:
    """Relative start second of every turn, plus total length in seconds."""
    times: list[int] = []
    clock = 1.0 if conv.voice else 0.0
    previous: Turn | None = None
    for turn in conv.turns:
        if previous is not None:
            words = len(previous.text.split())
            if conv.voice:
                rate = BOT_WORDS_PER_SEC if previous.speaker == BOT else CUSTOMER_WORDS_PER_SEC
                clock += words / rate
                clock += rng.uniform(0.8, 3.0) if turn.speaker == BOT else rng.uniform(1.5, 5.0)
            else:
                clock += rng.uniform(1, 2.5) if turn.speaker == BOT else 3 + words * 0.7 + rng.uniform(2, 12)
            if turn.wait_before is not None:
                clock += turn.wait_before
        times.append(int(round(clock)))
        previous = turn
    last_words = len(conv.turns[-1].text.split())
    tail = last_words / BOT_WORDS_PER_SEC if conv.voice else 2
    if conv.turns[-1].tags and "timeout_no_goodbye" in conv.turns[-1].tags:
        tail += CHAT_TIMEOUT_SEC if not conv.voice else VOICE_SILENCE_SEC * 2
    return times, int(round(times[-1] + tail + 1))


def hms(seconds: int) -> str:
    return f"{seconds // 3600:02d}:{seconds % 3600 // 60:02d}:{seconds % 60:02d}"


def render_transcript(conv: Conversation, times: list[int]) -> str:
    return "\n".join(f"[{hms(t)}] {turn.speaker}: {turn.text}" for turn, t in zip(conv.turns, times))


def build_conversation(spec: Spec, base_seed: int) -> dict:
    """Build one conversation and its golden twin. Returns the row, answer-key entry and golden."""
    facts_rng = random.Random(base_seed * 1000 + spec.number)
    start = WINDOW_START + timedelta(seconds=facts_rng.randint(0, WINDOW_DAYS * 86400 - 3600))
    start = start.replace(hour=8 + start.hour % 13)
    facts = draw_facts(facts_rng, start)

    # Same text rng seed for the real and the golden run, so both share the customer's wording.
    text_seed = base_seed * 7919 + spec.number
    conv = Conversation(spec, facts, random.Random(text_seed), spec.bugs)
    conv.build()
    conv.apply_ending(spec.ending)
    golden = Conversation(spec, facts, random.Random(text_seed), [])
    golden.build()

    times, length = turn_times(conv, random.Random(text_seed + 1))
    end = start + timedelta(seconds=length)

    timed_out = conv.turns[-1].text.startswith(("Are you still there?", "I haven't heard back"))
    if conv.escalation:
        closed_by, outcome = "agent", "escalated"
    elif spec.ending == "timeout" or timed_out:
        closed_by, outcome = "timeout", "abandoned"
    elif spec.ending == "hangup":
        closed_by, outcome = "customer", "abandoned"
    else:
        closed_by = "customer" if not conv.voice or facts_rng.random() < 0.6 else "bot"
        outcome = "resolved"

    metrics: dict = {}
    if len(conv.turns) > 3:
        metrics = {
            "auth_result": conv.auth,
            "intent_confidence": facts.confidence,
            "fallback_count": conv.fallbacks
            + sum(1 for t in conv.turns if t.text.startswith("Sorry, I didn't catch that. Please say")),
            "sentiment_end": {"resolved": "positive", "escalated": "neutral", "abandoned": "negative"}[outcome],
            "csat_offered": outcome == "resolved",
            "lookup_ms": facts.lookup_ms,
        }
        if conv.voice:
            metrics["asr_language"] = "en"

    row = {
        "conversation_id": spec.conversation_id,
        "channel": spec.channel,
        "started_at": start.strftime("%Y-%m-%dT%H:%M:%SZ"),
        "ended_at": end.strftime("%Y-%m-%dT%H:%M:%SZ"),
        "duration_sec": length,
        "turns": len(conv.turns),
        "closed_by": closed_by,
        "escalated": "true" if conv.escalation else "false",
        "escalation_reason": conv.escalation[1] if conv.escalation else "",
        "queue": conv.escalation[0] if conv.escalation else "",
        "outcome": outcome,
        "flow_steps": ";".join(name for name, _ in conv.steps),
        "custom_metrics": json.dumps(metrics, separators=(",", ":")),
        "transcript": render_transcript(conv, times),
    }

    bugs = []
    first_utterance = None
    for index, turn in enumerate(conv.turns, start=1):
        if "first_utterance" in turn.tags:
            first_utterance = {"turn": index, "text": turn.text}
        for tag in turn.tags:
            if tag in BUG_INFO:
                severity, meaning = BUG_INFO[tag]
                bugs.append({"type": tag, "turn": index, "severity": severity, "evidence": evidence(tag, conv)})
    key = {
        "conversation_id": spec.conversation_id,
        "transcript_file": f"transcripts/{spec.core_file}.txt" if spec.core_file else None,
        "channel": spec.channel,
        "expected_intent": spec.intent,
        "first_customer_utterance": first_utterance,
        "outcome": outcome,
        "bugs": bugs,
    }
    golden_doc = {
        "about": "What this conversation should have been. Invented data.",
        "conversation_id": spec.conversation_id,
        "channel": spec.channel,
        "expected_intent": spec.intent,
        "expected_outcome": "escalated" if golden.escalation else "resolved",
        "flow_steps": [name for name, _ in golden.steps],
        "turns": [{"speaker": t.speaker, "text": t.text} for t in golden.turns],
    }
    return {"row": row, "key": key, "golden": golden_doc, "spec": spec}


def evidence(tag: str, conv: Conversation) -> str:
    f = conv.f
    intent = conv.spec.intent
    if tag == "misroute":
        return f"Customer asked for {intent}; the bot routed to {MISROUTES[intent][0]}."
    if tag == "fallback_loop":
        return "Three fallback prompts in a row, then a transfer."
    if tag == "premature_escalation":
        return f"Transferred right after the request, without trying the {intent} flow."
    if tag == "timeout_no_goodbye":
        return "Last turn is 'Are you still there?'; the conversation then closed by timeout."
    if tag == "bot_repeats":
        return "The bot sent the same answer twice in a row."
    if tag == "wrong_order_id":
        return f"Customer gave {f.order_id}; the bot used {conv.used_order_id}."
    if tag == "ignored_correction":
        if intent == "change_delivery_address":
            return f"Customer corrected the address to {f.address2}; the bot kept {f.address}."
        if intent == "return_damaged_item":
            return f"Customer corrected the item to the {f.item}; the bot returned the {f.item2}."
        return "Customer pushed back on the wrong route; the bot fell back instead of switching."
    return ""


def variation_specs(rng: random.Random) -> list[Spec]:
    bag = list(VARIATION_BUG_BAG)
    rng.shuffle(bag)
    chat_slots = set(rng.sample(range(TOTAL_VARIATIONS), VARIATION_CHAT_COUNT))
    specs = []
    for offset, bug in enumerate(bag[:TOTAL_VARIATIONS]):
        number = len(CORE_SPECS) + offset + 1
        intent = rng.choice(BUG_INTENTS[bug]) if bug else rng.choice(list(INTENTS))
        channel = "chat" if offset in chat_slots else "voice"
        bugs = [bug] if bug else []
        ending = "resolved"
        if bug == "timeout_no_goodbye":
            ending = "timeout"
        elif bug is None and intent != "talk_to_human":
            roll = rng.random()
            if channel == "chat" and roll < 0.6:
                ending = "timeout"
            elif roll < 0.15:
                ending = "hangup"
        specs.append(Spec(number, intent, channel, bugs, ending))
    return specs


# -- writers ----------------------------------------------------------------------

COLUMNS = ["conversation_id", "channel", "started_at", "ended_at", "duration_sec", "turns", "closed_by",
           "escalated", "escalation_reason", "queue", "outcome", "flow_steps", "custom_metrics", "transcript"]


def write_text(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(text.encode("utf-8"))


def write_csv(path: Path, rows: list[dict]) -> None:
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=COLUMNS, lineterminator="\r\n")
    writer.writeheader()
    writer.writerows(rows)
    write_text(path, buffer.getvalue())


def _style_header(sheet) -> None:
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="F4F0FB")
        cell.fill = PatternFill("solid", fgColor="3A2A5A")
    sheet.freeze_panes = "A2"


def save_workbook(workbook: Workbook, path: Path) -> None:
    fixed = datetime(*ARCHIVE_TIME)
    workbook.properties.creator = "Kitsune Labs sample pack"
    workbook.properties.created = fixed
    workbook.properties.modified = fixed
    buffer = io.BytesIO()
    workbook.save(buffer)
    write_text_bytes(path, normalise_archive(buffer.getvalue()))


def normalise_archive(data: bytes) -> bytes:
    """Rewrite an .xlsx (a zip) with fixed timestamps so the bytes never change between runs."""
    source = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            content = source.read(info.filename)
            if info.filename == "docProps/core.xml":
                content = re.sub(rb"(<dcterms:(created|modified)[^>]*>)[^<]*", rb"\g<1>2025-03-14T00:00:00Z",
                                 content)
            entry = zipfile.ZipInfo(info.filename, date_time=ARCHIVE_TIME)
            entry.compress_type = zipfile.ZIP_DEFLATED
            entry.external_attr = 0o644 << 16
            target.writestr(entry, content)
    return out.getvalue()


def write_text_bytes(path: Path, data: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)


def write_conversations_xlsx(path: Path, rows: list[dict]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "conversations"
    sheet.append(COLUMNS)
    for row in rows:
        sheet.append([row[c] for c in COLUMNS])
    _style_header(sheet)
    widths = {"A": 14, "B": 8, "C": 21, "D": 21, "L": 40, "M": 40, "N": 90}
    for letter, width in widths.items():
        sheet.column_dimensions[letter].width = width
    for cell in sheet["N"][1:]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    save_workbook(workbook, path)


def write_intents_xlsx(path: Path) -> None:
    examples = {
        "cancel_order": ["cancel my order", "I don't want my order anymore", "please cancel MO-2210",
                         "can I stop an order before it ships", "undo my last purchase"],
        "refund_status": ["where is my refund", "I haven't got my money back", "refund for MO-3391",
                          "when will my refund arrive", "did you send my refund yet"],
        "change_delivery_address": ["change my delivery address", "send it to my work instead",
                                    "I moved, update the address", "wrong address on my order",
                                    "can I redirect my parcel"],
        "billing_dispute": ["I was charged twice", "there's a charge I don't recognise", "you billed me too much",
                            "double charge on my card", "why was I charged for MO-5120"],
        "track_parcel": ["where is my package", "track my order", "when will my parcel arrive",
                         "my order hasn't come yet", "tracking for MO-7731"],
        "gift_card_balance": ["gift card balance", "how much is on my gift card", "check my gift card",
                              "what's left on my card from Mousy", "is my gift card still valid"],
        "return_damaged_item": ["my kettle arrived broken", "return a damaged item", "the box was crushed",
                                "item came cracked", "I want a replacement for a broken item"],
        "talk_to_human": ["agent", "talk to a person", "representative please", "I want a human",
                          "connect me to support"],
        "update_payment_method": ["update my card", "change payment method", "my card expired",
                                  "add a new card", "remove my old card"],
        "product_question": ["is the toaster dishwasher safe", "does the lamp have a warranty",
                             "what's the cutting board made of", "how big are the mixing bowls",
                             "does the frother need batteries"],
    }
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "intents"
    sheet.append(["intent", "description", "example_1", "example_2", "example_3", "example_4", "example_5"])
    for intent, description in INTENTS.items():
        sheet.append([intent, description, *examples[intent]])
    _style_header(sheet)
    sheet.column_dimensions["A"].width = 26
    sheet.column_dimensions["B"].width = 60
    for letter in "CDEFG":
        sheet.column_dimensions[letter].width = 34
    save_workbook(workbook, path)


README_TEXT = """Mousy sample pack - Kitsune Labs
================================

ALL DATA IS INVENTED. Mousy is not a real company. Every customer, order, phone number
(555-01xx, a fictional range), address and conversation was written for this demo.

Mousy is an imaginary online shop (home, kitchen, small gadgets). Its support bot is
"Mousy Assist". The conversations come from voice calls and web chats.

Files
-----
mousy_conversations.csv    40 conversations, one per row. The full transcript is in one cell.
mousy_conversations.xlsx   The same 40 conversations as an Excel workbook.
mousy_batch_20.xlsx        The first 20 conversations, for bulk and pipeline demos.
transcripts/*.txt          8 single conversations, one per core scenario.
intents_v3.xlsx            10 intents with a description and 5 example utterances each.
golden/*.json              What each core conversation should have been (for the diff view).
answer_key.json            The bugs planted in each conversation, with turn numbers.

Columns
-------
conversation_id, channel (voice|chat), started_at, ended_at (UTC), duration_sec, turns,
closed_by (customer|bot|agent|timeout), escalated (true|false), escalation_reason, queue,
outcome (resolved|escalated|abandoned), flow_steps (semicolon list), custom_metrics (JSON),
transcript

Transcript format
-----------------
One line per turn:  [hh:mm:ss] Speaker: text
The time is relative to the start of the conversation.

Planted bug types
-----------------
misroute, fallback_loop, premature_escalation, timeout_no_goodbye, bot_repeats,
wrong_order_id, ignored_correction. See answer_key.json for what each one means.
"""

FILE_INFO = [
    ("mousy_conversations.csv", "Mousy conversations (CSV)",
     "40 conversations, one per row, transcript in one cell.",
     ["Test Case Generator", "Metrics Extractor", "First Utterance Finder"]),
    ("mousy_conversations.xlsx", "Mousy conversations (Excel)", "The same 40 conversations as a workbook.",
     ["Test Case Generator", "Metrics Extractor", "First Utterance Finder"]),
    ("mousy_batch_20.xlsx", "Batch of 20 (Excel)", "A 20-row subset for bulk runs.",
     ["Test Case Generator", "Report Builder"]),
    ("intents_v3.xlsx", "Intent list (Excel)", "10 intents, each with a description and 5 example utterances.",
     ["Utterance Generator"]),
    ("answer_key.json", "Answer key", "Every planted bug, with its turn number and severity.",
     ["Bug Finder", "Report Builder"]),
    ("README.txt", "Read me", "What each file is. All data is invented.", []),
]


def build_pack(out: Path) -> list[dict]:
    if out.exists():
        shutil.rmtree(out)
    out.mkdir(parents=True)

    rng = random.Random(SEED)
    specs = CORE_SPECS + variation_specs(rng)
    built = [build_conversation(spec, SEED) for spec in specs]
    rows = [b["row"] for b in built]

    write_csv(out / "mousy_conversations.csv", rows)
    write_conversations_xlsx(out / "mousy_conversations.xlsx", rows)
    write_conversations_xlsx(out / "mousy_batch_20.xlsx", rows[:20])
    write_intents_xlsx(out / "intents_v3.xlsx")
    for b in built:
        spec: Spec = b["spec"]
        if spec.core_file:
            write_text(out / "transcripts" / f"{spec.core_file}.txt", b["row"]["transcript"] + "\n")
            write_text(out / "golden" / f"{spec.core_file}.json", json.dumps(b["golden"], indent=2) + "\n")
    answer_key = {
        "about": "Planted bugs in the Mousy sample pack. All data is invented.",
        "bug_types": {name: {"severity": s, "meaning": m} for name, (s, m) in BUG_INFO.items()},
        "conversations": [b["key"] for b in built],
    }
    write_text(out / "answer_key.json", json.dumps(answer_key, indent=2) + "\n")
    write_text(out / "README.txt", README_TEXT)

    pack_files = sorted(p for p in out.rglob("*") if p.is_file())
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for path in pack_files:
            entry = zipfile.ZipInfo(f"mousy_sample_pack/{path.relative_to(out).as_posix()}", date_time=ARCHIVE_TIME)
            entry.compress_type = zipfile.ZIP_DEFLATED
            entry.external_attr = 0o644 << 16
            archive.writestr(entry, path.read_bytes())
    write_text_bytes(out / "mousy_sample_pack.zip", zip_buffer.getvalue())

    # Manifest for the Sample data page (not part of the zip: it describes the zip).
    entries = [{"path": "mousy_sample_pack.zip", "title": "Everything (.zip)",
                "purpose": "Every file below in one download.", "bytes": (out / "mousy_sample_pack.zip").stat().st_size,
                "tools": []}]
    for name, title, purpose, tools in FILE_INFO:
        entries.append({"path": name, "title": title, "purpose": purpose,
                        "bytes": (out / name).stat().st_size, "tools": tools})
    for folder, ext, kind, tools in (("transcripts", "txt", "Transcript", ["Bug Finder", "Transcript Converter"]),
                                     ("golden", "json", "Golden conversation", ["Bug Finder"])):
        for spec in CORE_SPECS:
            path = out / folder / f"{spec.core_file}.{ext}"
            label = spec.intent.replace("_", " ")
            bug_text = ", ".join(b.replace("_", " ") for b in spec.bugs) or "no planted bugs"
            purpose = (f"{spec.channel.title()} conversation: {label} ({bug_text})." if folder == "transcripts"
                       else f"What the {label} conversation should have been.")
            entries.append({"path": f"{folder}/{path.name}", "title": f"{kind}: {label}", "purpose": purpose,
                            "bytes": path.stat().st_size, "tools": tools})
    manifest = {"about": "Files in the Mousy sample pack. All data is invented.", "files": entries}
    write_text(out / "manifest.json", json.dumps(manifest, indent=2) + "\n")
    return rows


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = parser.parse_args()
    rows = build_pack(args.out)
    voice = sum(1 for r in rows if r["channel"] == "voice")
    escalated = sum(1 for r in rows if r["escalated"] == "true")
    abandoned = sum(1 for r in rows if r["outcome"] == "abandoned")
    print(f"Wrote {len(rows)} conversations to {args.out}")
    print(f"  voice {voice} / chat {len(rows) - voice} | escalated {escalated} | abandoned {abandoned}")


if __name__ == "__main__":
    main()
